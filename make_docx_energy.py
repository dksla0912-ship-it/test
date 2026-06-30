# -*- coding: utf-8 -*-
"""중3 과학 모의고사(Ⅵ. 에너지 전환과 보존) Word(.docx) 생성
- 2단(2열) 레이아웃, 배점 없음, 그림 삽입 (기존 make_docx.py와 동일 양식)
"""
import os
from PIL import Image
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from exam_data_energy import QUESTIONS, SECTIONS, ANSWERS, EXPLANATIONS, FIGS, CONFIG

FONT = "Malgun Gothic"
IMG_DIR = "images"


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi'):
        rf.set(qn(a), FONT)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def borders(table, color="999999", sz="4"):
    el = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        el.append(e)
    table._tbl.tblPr.append(el)


def set_columns(section, num=2, space_twips=360):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'), str(num))
    cols.set(qn('w:space'), str(space_twips))
    cols.set(qn('w:sep'), '1')


def para(doc, text="", size=10, bold=False, align=None, after=3, before=0, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text); set_font(r, size, bold, color)
    return p


def add_image(doc, fname, max_w_cm=7.2, max_h_cm=5.2):
    path = os.path.join(IMG_DIR, fname)
    with Image.open(path) as im:
        w, h = im.size
    ar = w / h
    wcm = max_w_cm; hcm = wcm / ar
    if hcm > max_h_cm:
        hcm = max_h_cm; wcm = hcm * ar
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(path, width=Cm(wcm))


def add_box(doc, lines, fill="F4F4F4", size=9):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, "B0B0B0", "4")
    c = t.cell(0, 0); shade(c, fill)
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        r = p.add_run(ln); set_font(r, size)
    para(doc, "", after=1)


def add_bogi(doc, items):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, "9DB3CC", "4")
    c = t.cell(0, 0); shade(c, "EEF3FA")
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("< 보 기 >"); set_font(r, 9, bold=True)
    for it in items:
        pp = c.add_paragraph(); pp.paragraph_format.space_after = Pt(0)
        rr = pp.add_run(it); set_font(rr, 9)
    para(doc, "", after=1)


def add_table(doc, headers, rows):
    n = len(headers)
    t = doc.add_table(rows=1 + len(rows), cols=n); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, h in enumerate(headers):
        cc = t.rows[0].cells[i]; cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cc.paragraphs[0].add_run(h); set_font(r, 8.5, bold=True); shade(cc, "E6E6E6")
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            cc = t.rows[ri+1].cells[ci]; cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cc.paragraphs[0].add_run(v); set_font(r, 8.5)
    para(doc, "", after=1)


def add_question(doc, q):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{q['num']}. "); set_font(r, 10, bold=True)
    r2 = p.add_run(q['stem']); set_font(r2, 10)

    has_fig = q['num'] in FIGS
    if has_fig:
        add_image(doc, FIGS[q['num']])
    for block in q.get('blocks', []):
        if block[0] == 'box':
            if has_fig:
                continue
            add_box(doc, block[1])
        elif block[0] == 'table':
            add_table(doc, block[1], block[2])
        elif block[0] == 'bogi':
            add_bogi(doc, block[1])

    for ch in q['choices']:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.left_indent = Cm(0.3)
        rr = cp.add_run(ch); set_font(rr, 9.5)


def section_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); set_font(r, 12, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))


def header_block(doc):
    t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, "333333", "8")
    left, mid, right = t.rows[0].cells
    lp = left.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(lp.add_run(CONFIG["grade_label"]), 10, bold=True)
    mp = mid.paragraphs[0]; mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(mp.add_run(CONFIG["exam_title"]), 15, bold=True)
    rp = right.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(rp.add_run(CONFIG["range_label_docx"]), 8.5)
    para(doc, "", after=2)
    for ln in CONFIG["note_lines"]:
        para(doc, ln, size=8.5, color=RGBColor(0x60, 0x60, 0x60), after=1)


def base_doc():
    doc = Document()
    s = doc.sections[0]
    s.left_margin = Cm(1.3); s.right_margin = Cm(1.3)
    s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.3)
    st = doc.styles['Normal']
    st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return doc


def build_exam():
    doc = base_doc()
    header_block(doc)
    doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(doc.sections[0], 1)
    set_columns(doc.sections[1], 2)

    for q in QUESTIONS:
        if q['num'] in SECTIONS:
            section_title(doc, SECTIONS[q['num']])
        add_question(doc, q)
    para(doc, "", after=4)
    para(doc, "― 문제 끝 ―", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=10)
    out = CONFIG["out_exam"] + ".docx"
    doc.save(out)
    print("저장:", out)


def build_answers():
    doc = base_doc()
    para(doc, CONFIG["answer_title"], size=14, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, "빠른 정답표", size=12, bold=True, after=3, color=RGBColor(0x1F, 0x3A, 0x5F))
    headers = ["문항", "정답", "문항", "정답", "문항", "정답", "문항", "정답"]
    rows = []
    for i in range(10):
        row = []
        for j in range(4):
            num = i + 1 + j*10
            row += [str(num), ANSWERS[num]]
        rows.append(row)
    add_table(doc, headers, rows)
    para(doc, "", after=4)

    for n in range(1, 41):
        if n in SECTIONS:
            section_title(doc, SECTIONS[n])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{n}. 정답 {ANSWERS[n]}")
        set_font(r, 10.5, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))
        ep = doc.add_paragraph(); ep.paragraph_format.space_after = Pt(2)
        er = ep.add_run(EXPLANATIONS[n]); set_font(er, 9.5)
    out = CONFIG["out_answer"] + ".docx"
    doc.save(out)
    print("저장:", out)


if __name__ == "__main__":
    build_exam()
    build_answers()

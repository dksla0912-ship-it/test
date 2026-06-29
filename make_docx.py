# -*- coding: utf-8 -*-
"""중2 과학 모의고사 시험지 / 정답해설 Word(.docx) 생성 스크립트"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Malgun Gothic"  # 맑은 고딕


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)
    rfonts.set(qn('w:ascii'), FONT)
    rfonts.set(qn('w:hAnsi'), FONT)


def shade_cell(cell, fill="F2F2F2"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def set_table_borders(table, color="888888", sz="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def add_para(doc, text="", size=10.5, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    return p


def add_box(doc, lines, fill="F7F7F7", border=True, size=10):
    """설명/자료 글상자: 단일 셀 표"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if border:
        set_table_borders(table, color="AAAAAA", sz="4")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.paragraphs[0].text = ""
    first = True
    for line in lines:
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        r = p.add_run(line)
        set_font(r, size=size)
    add_para(doc, "", space_after=2)
    return table


def add_data_table(doc, headers, rows, widths=None):
    n = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = hdr[i].paragraphs[0].add_run(h)
        set_font(r, size=9.5, bold=True)
        shade_cell(hdr[i], "E8E8E8")
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cells[ci].paragraphs[0].add_run(val)
            set_font(r, size=9.5)
    add_para(doc, "", space_after=2)
    return table


def add_question(doc, q):
    # 문항 번호 + 발문
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{q['num']}. ")
    set_font(r, size=10.5, bold=True)
    r2 = p.add_run(q['stem'])
    set_font(r2, size=10.5, bold=False)
    r3 = p.add_run(f"  [{q.get('pt','2.5점')}]")
    set_font(r3, size=9, bold=False, color=RGBColor(0x70, 0x70, 0x70))

    for block in q.get('blocks', []):
        if block[0] == 'box':
            add_box(doc, block[1])
        elif block[0] == 'table':
            add_data_table(doc, block[1], block[2])
        elif block[0] == 'bogi':
            add_box(doc, ["<보기>"] + block[1], fill="EFF4FB")

    for ch in q['choices']:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(1)
        cp.paragraph_format.left_indent = Cm(0.4)
        rr = cp.add_run(ch)
        set_font(rr, size=10.5)


def add_title_block(doc, subtitle):
    p = add_para(doc, "2026학년도 중학교 2학년 과학 모의고사", size=16, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, subtitle, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2,
             color=RGBColor(0x55, 0x55, 0x55))
    # 구분선
    pp = doc.add_paragraph()
    pborder = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '333333')
    pborder.append(bottom)
    pp.paragraph_format.element.get_or_add_pPr().append(pborder)


def section(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=13, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))


from exam_data import QUESTIONS, SECTIONS, ANSWERS, EXPLANATIONS


def build_exam():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    add_title_block(doc, "출제 범위 : Ⅴ.식물과 에너지 · Ⅵ.동물과 에너지 · Ⅶ.전류의 자기 작용   |   40문항 (전 문항 5지선다)   |   60분   |   100점")
    add_para(doc, "※ 모든 문항은 5지선다형 객관식입니다. <보기>가 있는 문항은 옳은 것을 모두 고른 조합을 선택하세요.",
             size=9, color=RGBColor(0x70,0x70,0x70), space_after=6)
    for q in QUESTIONS:
        if q['num'] in SECTIONS:
            section(doc, SECTIONS[q['num']])
        add_question(doc, q)
    add_para(doc, "", space_after=6)
    add_para(doc, "— 문제 끝 —", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11)
    doc.save("2026_중2_과학_모의고사_문제.docx")
    print("저장: 2026_중2_과학_모의고사_문제.docx")


def build_answers():
    doc = Document()
    for s in doc.sections:
        s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.6)
    add_title_block(doc, "정답 및 해설")
    add_para(doc, "빠른 정답표", size=13, bold=True, space_after=4, color=RGBColor(0x1F,0x3A,0x5F))
    # 10 x (문항/정답 4쌍)
    headers = ["문항","정답","문항","정답","문항","정답","문항","정답"]
    rows = []
    for i in range(10):
        row = []
        for j in range(4):
            num = i + 1 + j*10
            row += [str(num), ANSWERS[num]]
        rows.append(row)
    add_data_table(doc, headers, rows)

    cur = None
    for n in range(1, 41):
        if n in SECTIONS:
            section(doc, SECTIONS[n])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{n}. 정답 {ANSWERS[n]}")
        set_font(r, size=10.5, bold=True, color=RGBColor(0xC0,0x39,0x2B))
        ep = doc.add_paragraph()
        ep.paragraph_format.space_after = Pt(3)
        er = ep.add_run(EXPLANATIONS[n])
        set_font(er, size=10)
    doc.save("2026_중2_과학_모의고사_정답해설.docx")
    print("저장: 2026_중2_과학_모의고사_정답해설.docx")


if __name__ == "__main__":
    build_exam()
    build_answers()
